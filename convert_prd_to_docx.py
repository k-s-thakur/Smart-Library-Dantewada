import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>')
    tcPr.append(shd)


def create_styled_prd_docx():
    md_path = r"E:\KHAGENDRA SINGH THAKUR\DEVOPS BUILD\SMART LIBRARY APP\PRD_SMART_LIBRARY_SYSTEM.md"
    docx_path = r"E:\KHAGENDRA SINGH THAKUR\DEVOPS BUILD\SMART LIBRARY APP\PRD_SMART_LIBRARY_SYSTEM.docx"

    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = Document()

    # Set 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    lines = md_text.split("\n")
    i = 0
    in_code_block = False
    code_content = []
    in_table = False
    table_rows = []

    def flush_table(rows):
        if not rows:
            return
        # Filter out markdown divider lines |---|---|
        clean_rows = []
        for r in rows:
            if re.match(r"^\|?[\s\-:]+(\|[\s\-:]+)+\|?$", r.strip()):
                continue
            cells = [c.strip() for c in r.strip().strip("|").split("|")]
            clean_rows.append(cells)

        if not clean_rows:
            return

        cols_count = max(len(r) for r in clean_rows)
        # Pad shorter rows
        for r in clean_rows:
            while len(r) < cols_count:
                r.append("")

        table = doc.add_table(rows=len(clean_rows), cols=cols_count)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        for row_idx, r in enumerate(clean_rows):
            is_header = (row_idx == 0)
            for col_idx, cell_value in enumerate(r):
                cell = table.cell(row_idx, col_idx)
                cell.text = cell_value
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)

                if is_header:
                    set_cell_background(cell, "1E3A8A")
                    for run in p.runs:
                        run.font.bold = True
                        run.font.name = "Arial"
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                    set_cell_background(cell, bg)
                    for run in p.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(51, 65, 85)

        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_after = Pt(6)

    def flush_code(lines):
        if not lines:
            return
        code_str = "\n".join(lines)
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        p = cell.paragraphs[0]
        p.text = code_str
        for run in p.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(226, 232, 240)
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_after = Pt(6)

    while i < len(lines):
        line = lines[i]

        # Code block handler
        if line.strip().startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code(code_content)
                code_content = []
            else:
                if in_table:
                    in_table = False
                    flush_table(table_rows)
                    table_rows = []
                in_code_block = True
                code_content = []
            i += 1
            continue

        if in_code_block:
            code_content.append(line)
            i += 1
            continue

        # Table rows
        if line.strip().startswith("|") and line.strip().endswith("|"):
            in_table = True
            table_rows.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                flush_table(table_rows)
                table_rows = []

        stripped = line.strip()

        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=0)
            run = h.add_run(stripped[2:])
            run.font.name = "Arial"
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
        elif stripped.startswith("## "):
            h = doc.add_heading(level=1)
            run = h.add_run(stripped[3:])
            run.font.name = "Arial"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
        elif stripped.startswith("### "):
            h = doc.add_heading(level=2)
            run = h.add_run(stripped[4:])
            run.font.name = "Arial"
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
        elif stripped.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
        elif stripped.startswith("* ") or stripped.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            text = stripped[2:]
            # parse bold
            parts = re.split(r"(\*\*.*?\*\*)", text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(2)
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style='List Number')
            text = re.sub(r"^\d+\.\s", "", stripped)
            parts = re.split(r"(\*\*.*?\*\*)", text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
            p.paragraph_format.space_after = Pt(2)
        elif stripped != "":
            p = doc.add_paragraph()
            parts = re.split(r"(\*\*.*?\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(30, 41, 59)
            p.paragraph_format.space_after = Pt(4)

        i += 1

    if in_table:
        flush_table(table_rows)

    doc.save(docx_path)
    print("SUCCESS: Word document created at", docx_path)

if __name__ == "__main__":
    create_styled_prd_docx()
